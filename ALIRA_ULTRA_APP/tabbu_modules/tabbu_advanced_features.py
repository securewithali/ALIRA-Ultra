# -*- coding: utf-8 -*-
"""TITAN Advanced Features v2.0 — Complete Rewrite"""

from alira_slot import slot as _ALIRA_slot
import os, sys, re, json, time, base64, threading, platform
import subprocess, tempfile, io, datetime, glob
from pathlib import Path
from typing import Optional, List, Dict, Any

try:
    import cv2
    import numpy as np
    CV2_OK = True
except ImportError:
    CV2_OK = False; np = None

try:
    import pyautogui
    pyautogui.FAILSAFE = True; pyautogui.PAUSE = 0.05
    PYAUTOGUI_OK = True
except ImportError:
    PYAUTOGUI_OK = False

try:
    import pyperclip; CLIP_OK = True
except ImportError:
    CLIP_OK = False

try:
    from PIL import Image as PILImage; PIL_OK = True
except ImportError:
    PIL_OK = False

try:
    import pytesseract
    if platform.system() == "Windows":
        for tp in [r"C:\Program Files\Tesseract-OCR\tesseract.exe",
                   r"C:\Program Files (x86)\Tesseract-OCR\tesseract.exe"]:
            if os.path.exists(tp):
                pytesseract.pytesseract.tesseract_cmd = tp; break
    OCR_OK = True
except ImportError:
    OCR_OK = False

try:
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib.units import cm
    from reportlab.lib import colors
    from reportlab.platypus import (SimpleDocTemplate, Paragraph, Spacer,
                                    PageBreak, HRFlowable)
    from reportlab.lib.enums import TA_LEFT, TA_CENTER, TA_JUSTIFY
    REPORTLAB_OK = True
except ImportError:
    REPORTLAB_OK = False

try:
    import pypdf
    PYPDF_OK = True
except ImportError:
    PYPDF_OK = False

try:
    import docx
    DOCX_OK = True
except ImportError:
    DOCX_OK = False

try:
    import mss
    MSS_OK = True
except ImportError:
    MSS_OK = False

try:
    import win32com.client
    WIN32COM_OK = True
except ImportError:
    WIN32COM_OK = False

from PyQt5.QtWidgets import (
    QWidget, QVBoxLayout, QHBoxLayout, QLabel, QPushButton,
    QTextEdit, QFileDialog, QFrame, QDialog, QApplication,
    QMessageBox, QSizePolicy, QGridLayout, QLineEdit
)
from PyQt5.QtCore import Qt, QTimer, QThread, pyqtSignal, QRect
from PyQt5.QtGui import QColor, QFont

def _get_root() -> Path:
    if getattr(sys, 'frozen', False):
        return Path(sys.executable).parent
    p = Path(os.path.abspath(__file__)).parent
    return p.parent if p.name == "TITAN_modules" else p

ROOT_PATH  = _get_root()
MEM_DIR    = ROOT_PATH / "TITAN_memory"
FACE_DIR   = ROOT_PATH / "TITAN_face_data"
CAM_DIR    = ROOT_PATH / "TITAN_webcam"
ASSIGN_DIR = ROOT_PATH / "TITAN_assignments"
for d in [MEM_DIR, FACE_DIR, CAM_DIR, ASSIGN_DIR]:
    d.mkdir(parents=True, exist_ok=True)

MEM_FILE      = MEM_DIR / "life_memory.json"
PROFILES_FILE = FACE_DIR / "user_profiles.json"

BG="#050810"; ACCENT="#e8406a"; ACCENT2="#ff6b9d"
CYAN="#00d4ff"; GOLD="#ffd060"; TEXT="#e8eaf6"
MUTED="#6b7394"; GREEN="#2ecc71"; BG3="#0d1525"

_API_KEY = None

def _analyze_image_with_gemini(image_path: str, prompt: str) -> str:
    if not _API_KEY:
        print("[Vision] Error: API Key not set")
        return "API key not set. AI model cannot analyze image without key."
    try:
        import urllib.request
        with open(image_path, "rb") as f:
            img_data = f.read()

        ext  = Path(image_path).suffix.lower().lstrip(".")
        mime = {"jpg":"image/jpeg","jpeg":"image/jpeg","png":"image/png",
                "bmp":"image/bmp","gif":"image/gif","webp":"image/webp"}.get(ext,"image/jpeg")

        # ✅ FIX: bade PNG screenshot (1+ MB) upload pe "read operation timed out"
        # aa raha tha. Max 1600px + JPEG q85 mein compress karke bhejo (~10x chhota).
        if len(img_data) > 300 * 1024:
            try:
                import io
                from PIL import Image
                im = Image.open(io.BytesIO(img_data)).convert("RGB")
                im.thumbnail((1600, 1600))
                buf = io.BytesIO(); im.save(buf, "JPEG", quality=85)
                img_data, mime = buf.getvalue(), "image/jpeg"
            except Exception as ce:
                print(f"[Vision] compress skipped: {ce}")
        img_b64 = base64.b64encode(img_data).decode()

        payload = json.dumps({
            "contents": [{
                "parts": [
                    {"text": prompt},
                    {"inline_data": {"mime_type": mime, "data": img_b64}}
                ]
            }],
            "generationConfig": {
                "temperature": 0.4,
                "topK": 32,
                "topP": 1,
                "maxOutputTokens": 2048,
            }
        }).encode()

        # ✅ FIX: gemini-1.5-flash retired hai (404 "not found"/"no longer
        # available") — yehi "[Vision] Exception: HTTP Error 404" ka karan
        # tha jo har screenshot/camera analyze call pe fail hota tha.
        print(f"[Vision] Sending to Gemini: {Path(image_path).name} ({len(img_data)//1024} KB)")
        # ✅ FIX: gemini-2.5-flash free quota (429) khatam hone par model fallback chain
        from TABBU_modules.gemini_rest import generate as _gen
        res_data, _used = _gen(json.loads(payload), _API_KEY, timeout=60)
        if True:
            cands = res_data.get("candidates", [])
            if cands:
                parts = cands[0].get("content", {}).get("parts", [])
                text = " ".join(p.get("text", "") for p in parts)
                if text.strip():
                    print("[Vision] Analysis success")
                    return text.strip()

            error_msg = res_data.get("error", {}).get("message", "Unknown API error")
            print(f"[Vision] API Error: {error_msg}")
            return f"Vision API Error: {error_msg}"

    except Exception as e:
        print(f"[Vision] Exception: {e}")
        return f"Vision system error: {e}"

def extract_text_from_image(image_path: str) -> str:
    if not os.path.exists(image_path):
        return f"Image not found: {image_path}"
    if _API_KEY:
        r = _analyze_image_with_gemini(image_path,
            "Extract ALL text from this image exactly as written. "
            "If it is a question paper, list every question numbered. "
            "Return only extracted text, no commentary.")
        if r and "error" not in r.lower()[:20]:
            return r
    if OCR_OK and PIL_OK:
        try:
            img  = PILImage.open(image_path).convert("L")
            text = pytesseract.image_to_string(img)
            if text.strip():
                return text.strip()
        except Exception as e:
            print(f"[OCR] {e}")
    return "OCR failed. Install: pip install pytesseract pillow"

class FileBrowser:
    @classmethod
    def _roots(cls):
        roots = []
        if platform.system()=="Windows":
            import string
            for d in string.ascii_uppercase:
                if os.path.exists(f"{d}:\\"):
                    roots.append(f"{d}:\\")
        home = Path.home()
        for sub in ["Desktop","Documents","Downloads","Pictures","Videos"]:
            p = home/sub
            if p.exists(): roots.insert(0,str(p))
        roots.insert(0, str(Path.cwd()))
        return roots

    @classmethod
    def find_file(cls, name: str, max_results: int=10) -> list:
        found = []
        for root in cls._roots():
            try:
                for dp,dns,fns in os.walk(root):
                    dns[:] = [d for d in dns if not d.startswith('.')
                               and d not in {'Windows','System32','$Recycle.Bin',
                                             'node_modules','__pycache__'}]
                    for fn in fns:
                        if name.lower() in fn.lower():
                            found.append(os.path.join(dp,fn))
                            if len(found)>=max_results: return found
            except (PermissionError,OSError): continue
        return found

    @classmethod
    def list_folder(cls, path: str) -> str:
        try:
            p = Path(path)
            if not p.exists(): return f"Folder not found: {path}"
            items = sorted(p.iterdir())
            files   = [i for i in items if i.is_file() and not i.name.startswith('.')]
            folders = [i for i in items if i.is_dir()  and not i.name.startswith('.')]
            lines   = [f"📂 {path}"]
            if folders:
                lines.append(f"Folders ({len(folders)}):")
                for f in folders[:15]: lines.append(f"  📁 {f.name}")
            if files:
                lines.append(f"Files ({len(files)}):")
                for f in files[:25]:
                    sz = f.stat().st_size
                    lines.append(f"  📄 {f.name} ({sz//1024}KB)" if sz>1024
                                  else f"  📄 {f.name} ({sz}B)")
            return "\n".join(lines)
        except PermissionError:
            return f"Permission denied: {path}"

    @classmethod
    def recent_files(cls) -> str:
        home  = Path.home()
        lines = ["📁 Recent Files:"]
        for folder in [home/"Desktop", home/"Documents",
                       home/"Downloads", Path.cwd()]:
            if not folder.exists(): continue
            files = sorted(folder.glob("*"),
                           key=lambda f: f.stat().st_mtime if f.is_file() else 0,
                           reverse=True)[:5]
            if files:
                lines.append(f"\n📂 {folder.name}/")
                for f in files:
                    if f.is_file(): lines.append(f"  • {f.name}")
        return "\n".join(lines)

    @classmethod
    def find_pdfs(cls, max_results: int=20) -> list:
        found = []
        for root in cls._roots():
            try:
                for dp,dns,fns in os.walk(root):
                    dns[:] = [d for d in dns if not d.startswith('.')]
                    for fn in fns:
                        if fn.lower().endswith(".pdf"):
                            found.append(os.path.join(dp,fn))
                            if len(found)>=max_results: return found
            except (PermissionError,OSError): continue
        return found

class TITANMemory:
    _data: Dict = {}
    _lock = threading.Lock()

    @classmethod
    def load(cls):
        try:
            if MEM_FILE.exists():
                with open(MEM_FILE,"r",encoding="utf-8") as f:
                    cls._data = json.load(f)
        except Exception: cls._data = {}
        for k in ["conversations","notes","facts","events","preferences"]:
            cls._data.setdefault(k,[])

    @classmethod
    def save(cls):
        try:
            with cls._lock:
                with open(MEM_FILE,"w",encoding="utf-8") as f:
                    json.dump(cls._data,f,ensure_ascii=False,indent=2)
        except Exception as e: print(f"[Memory] {e}")

    @classmethod
    def add_conversation(cls, role: str, content: str, context: str=""):
        cls._data["conversations"].append({
            "role":role,"content":content,"context":context,
            "timestamp":datetime.datetime.now().isoformat()})
        if len(cls._data["conversations"])>500:
            cls._data["conversations"]=cls._data["conversations"][-500:]
        cls.save()

    @classmethod
    def add_note(cls, title: str, content: str, tags: list=None):
        cls._data["notes"].append({
            "id":len(cls._data["notes"])+1,"title":title,
            "content":content,"tags":tags or [],
            "created":datetime.datetime.now().isoformat()})
        cls.save()

    @classmethod
    def add_fact(cls, fact: str, source: str=""):
        cls._data["facts"].append({
            "fact":fact,"source":source,
            "timestamp":datetime.datetime.now().isoformat()})
        cls.save()

    @classmethod
    def search(cls, query: str, limit: int=10) -> list:
        results=[]; q=query.lower()
        for item in cls._data.get("conversations",[])[-200:]:
            if q in item.get("content","").lower():
                results.append({"type":"conversation",**item})
        for item in cls._data.get("notes",[]):
            if q in item.get("title","").lower() or q in item.get("content","").lower():
                results.append({"type":"note",**item})
        for item in cls._data.get("facts",[]):
            if q in item.get("fact","").lower():
                results.append({"type":"fact",**item})
        return results[:limit]

    @classmethod
    def get_recent_context(cls, n: int=15) -> str:
        recent=cls._data.get("conversations",[])[-n:]
        return "\n".join(
            f"{'User' if i['role']=='user' else 'TITAN'}: {i['content'][:200]}"
            for i in recent)

    @classmethod
    def get_summary(cls) -> dict:
        return {k:len(cls._data.get(k,[])) for k in ["conversations","notes","facts","events"]}

class DocReader:
    @staticmethod
    def extract(path: str) -> str:
        if not os.path.exists(path):
            return f"File not found: {path}"

        ext = Path(path).suffix.lower()

        # .docx Files
        if ext == ".docx":
            if not DOCX_OK:
                return "❌ python-docx not installed. Bol: 'pip install python-docx'"
            try:
                doc = docx.Document(path)
                full_text = []
                for para in doc.paragraphs:
                    full_text.append(para.text)
                for table in doc.tables:
                    for row in table.rows:
                        row_text = [cell.text for cell in row.cells]
                        full_text.append(" | ".join(row_text))
                return "\n".join(full_text).strip()
            except Exception as e:
                return f"❌ Error reading .docx file: {e}"

        # .doc Files (Legacy)
        elif ext == ".doc":
            if platform.system() != "Windows":
                return "❌ .doc reading only supported on Windows"
            if not WIN32COM_OK:
                return "❌ pywin32 not installed. Bol: 'pip install pywin32'"
            try:
                import win32com.client
                word = win32com.client.Dispatch("Word.Application")
                word.Visible = False
                doc = word.Documents.Open(os.path.abspath(path))
                text = doc.Content.Text
                doc.Close()
                word.Quit()
                return text.strip()
            except Exception as e:
                return f"❌ Error reading .doc file: {e}"

        return f"❌ Unsupported file type: {ext}"

    @staticmethod
    def find_and_extract(path_or_name: str) -> tuple:
        if os.path.exists(path_or_name):
            return path_or_name, DocReader.extract(path_or_name)

        # Word files dhoondein (.docx, .doc)
        found = FileBrowser.find_file(path_or_name, max_results=3)
        if found:
            # Filter for Word extensions if multiple files found
            word_files = [f for f in found if f.lower().endswith(('.docx', '.doc'))]
            target = word_files[0] if word_files else found[0]
            return target, DocReader.extract(target)

        return None, f"Not found: {path_or_name}"

class PDFReader:
    @staticmethod
    def extract(path: str) -> str:
        if not os.path.exists(path):
            return f"File not found: {path}"
        text=""
        try:
            import pdfplumber
            with pdfplumber.open(path) as pdf:
                for pg in pdf.pages:
                    t=pg.extract_text()
                    if t: text+=t+"\n\n"
            if text.strip(): return text.strip()
        except ImportError: pass
        except Exception as e: print(f"[PDF] pdfplumber: {e}")
        try:
            import fitz
            doc=fitz.open(path)
            for pg in doc: text+=pg.get_text()+"\n\n"
            if text.strip(): return text.strip()
        except ImportError: pass
        except Exception as e: print(f"[PDF] fitz: {e}")
        try:
            import PyPDF2
            with open(path,"rb") as f:
                r=PyPDF2.PdfReader(f)
                for pg in r.pages: text+=(pg.extract_text() or "")+"\n\n"
            if text.strip(): return text.strip()
        except ImportError: pass
        except Exception as e: print(f"[PDF] PyPDF2: {e}")
        if _API_KEY:
            try:
                import fitz
                doc=fitz.open(path)
                texts=[]
                for i,pg in enumerate(doc[:4]):
                    pix=pg.get_pixmap(dpi=150)
                    img_path=str(CAM_DIR/f"pdf_pg_{i}.png")
                    pix.save(img_path)
                    t=_analyze_image_with_gemini(img_path,
                        "Extract all text from this PDF page. Preserve structure.")
                    texts.append(t)
                return "\n\n".join(texts)
            except Exception: pass
        return "PDF read failed. Install: pip install pdfplumber"

    @staticmethod
    def find_and_extract(path_or_name: str) -> tuple:
        if os.path.exists(path_or_name):
            return path_or_name, PDFReader.extract(path_or_name)
        found=FileBrowser.find_file(path_or_name,max_results=3)
        if found:
            return found[0], PDFReader.extract(found[0])
        return None, f"Not found: {path_or_name}"

    @staticmethod
    def summarize(text: str, n: int=10) -> str:
        sents=[s.strip() for s in text.split(".") if len(s.strip())>30]
        if not sents: return text[:500]
        freq={}
        for w in text.lower().split():
            if len(w)>4: freq[w]=freq.get(w,0)+1
        scored=sorted([(sum(freq.get(w.lower(),0) for w in s.split()),s)
                       for s in sents],reverse=True)
        return ". ".join(s for _,s in scored[:n])+"."

    @staticmethod
    def quiz(text: str, n: int=5) -> list:
        import random
        sents=[s.strip() for s in text.split(".") if len(s.strip())>50]
        if not sents: return []
        sel=random.sample(sents,min(n,len(sents)))
        out=[]
        for i,s in enumerate(sel):
            ws=s.split()
            if len(ws)<8: continue
            ans=ws[len(ws)//2]
            out.append({"q":f"Q{i+1}: {s.replace(ans,'_____',1)}","answer":ans})
        return out

class WebcamThread(QThread):
    sig_ok    = pyqtSignal(str)
    sig_error = pyqtSignal(str)
    def __init__(self, save_dir: Path):
        super().__init__(); self._dir=save_dir
    def run(self):
        if not CV2_OK:
            self.sig_error.emit("OpenCV not installed (pip install opencv-python)"); return
        cap=None
        backends = ([cv2.CAP_DSHOW,cv2.CAP_MSMF,cv2.CAP_ANY]
                    if platform.system()=="Windows" else [cv2.CAP_ANY])
        for be in backends:
            for idx in range(3):
                try:
                    c=cv2.VideoCapture(idx,be)
                    if c.isOpened():
                        ret,fr=c.read()
                        if ret and fr is not None:
                            cap=c; break
                    c.release()
                except Exception: pass
            if cap: break
        if not cap:
            self.sig_error.emit("No webcam found / access denied"); return
        for _ in range(12): cap.read(); time.sleep(0.03)
        ret,frame=cap.read(); cap.release()
        if not ret or frame is None:
            self.sig_error.emit("Frame capture failed"); return
        self._dir.mkdir(parents=True,exist_ok=True)
        ts=datetime.datetime.now().strftime("%Y%m%d_%H%M%S")
        path=str(self._dir/f"cam_{ts}.jpg")
        if cv2.imwrite(path,frame,[cv2.IMWRITE_JPEG_QUALITY,92]):
            self.sig_ok.emit(path)
        else:
            self.sig_error.emit(f"Could not write {path}")

class ScreenCapture:
    @staticmethod
    def capture(save_path: str=None) -> Optional[str]:
        # Ensure CAM_DIR exists at project root level
        CAM_DIR.mkdir(parents=True, exist_ok=True)

        if not save_path:
            ts=datetime.datetime.now().strftime("%Y%m%d_%H%M%S")
            save_path=str(CAM_DIR/f"screen_{ts}.png")

        print(f"[Vision] Taking screenshot: {save_path}")

        # 1. Try mss (Fastest and most reliable on multi-monitor)
        if MSS_OK:
            try:
                with mss.mss() as sct:
                    sct.shot(output=save_path)
                    if os.path.exists(save_path):
                        print(f"[Vision] Screenshot success (mss)")
                        return save_path
            except Exception as e: print(f"[Vision] mss error: {e}")

        # 2. Try pyautogui
        if PYAUTOGUI_OK:
            try:
                import pyautogui
                img=pyautogui.screenshot()
                img.save(save_path)
                if os.path.exists(save_path):
                    print(f"[Vision] Screenshot success (PyAutoGUI)")
                    return save_path
            except Exception as e: print(f"[Vision] PyAutoGUI error: {e}")

        # 3. Try PIL ImageGrab
        if PIL_OK:
            try:
                from PIL import ImageGrab
                img=ImageGrab.grab()
                img.save(save_path)
                if os.path.exists(save_path):
                    print(f"[Vision] Screenshot success (PIL)")
                    return save_path
            except Exception as e: print(f"[Vision] PIL error: {e}")

        return None

class ClipboardAutomation:
    @staticmethod
    def get() -> str:
        if CLIP_OK:
            try: return pyperclip.paste()
            except Exception: pass
        if platform.system()=="Windows":
            try:
                r=subprocess.run(["powershell","-command","Get-Clipboard"],
                                 capture_output=True,text=True)
                return r.stdout.strip()
            except Exception: pass
        return ""

    @staticmethod
    def set(text: str) -> bool:
        if CLIP_OK:
            try: pyperclip.copy(text); return True
            except Exception: pass
        if platform.system()=="Windows":
            try:
                safe=text.replace("'","''")
                subprocess.run(["powershell","-command",
                                f"Set-Clipboard -Value '{safe}'"],
                               check=True,capture_output=True)
                return True
            except Exception: pass
        return False

    @staticmethod
    def copy_from_app(window_title: str) -> str:
        if not PYAUTOGUI_OK: return "pyautogui not installed"
        try:
            import pygetwindow as gw
            wins=gw.getWindowsWithTitle(window_title)
            if not wins: return f"Window '{window_title}' not found"
            wins[0].activate(); time.sleep(0.3)
            pyautogui.hotkey("ctrl","a"); time.sleep(0.1)
            pyautogui.hotkey("ctrl","c"); time.sleep(0.2)
            return ClipboardAutomation.get()
        except Exception as e: return f"Error: {e}"

    @staticmethod
    def paste_to_app(window_title: str, text: str) -> bool:
        if not PYAUTOGUI_OK: return False
        try:
            ClipboardAutomation.set(text)
            import pygetwindow as gw
            wins=gw.getWindowsWithTitle(window_title)
            if not wins: return False
            wins[0].activate(); time.sleep(0.3)
            pyautogui.hotkey("ctrl","v"); return True
        except Exception: return False

    @staticmethod
    def run_in_vscode(code: str) -> dict:
        if not PYAUTOGUI_OK: return {"success":False,"error":"pyautogui not installed"}
        try:
            ClipboardAutomation.set(code)
            import pygetwindow as gw
            wins=[w for w in gw.getAllWindows()
                  if "visual studio code" in w.title.lower() or "code" in w.title.lower()]
            if not wins: return {"success":False,"error":"VSCode not found"}
            wins[0].activate(); time.sleep(0.4)
            pyautogui.hotkey("ctrl","`"); time.sleep(0.6)
            pyautogui.hotkey("ctrl","v"); time.sleep(0.1)
            pyautogui.press("enter")
            return {"success":True,"message":"Code pasted and run in VSCode"}
        except Exception as e: return {"success":False,"error":str(e)}

class ErrorDetector:
    DANGEROUS=[r"rm\s+-rf",r"del\s+/[sfq]",r"format\s+[a-z]:",
               r"DROP\s+DATABASE",r"shutil\.rmtree",r"dd\s+if="]
    FIXES={
        "ModuleNotFoundError":"pip install {module}",
        "ImportError":"pip install {module}",
        "SyntaxError":"Line {line} par syntax error — brackets/quotes check karo",
        "IndentationError":"Indentation galat — spaces ya tabs consistent rakho",
        "NameError":"'{name}' define nahi — pehle declare karo",
        "TypeError":"Type mismatch — function arguments check karo",
        "AttributeError":"Object mein ye attribute nahi — spelling check karo",
        "FileNotFoundError":"File '{file}' nahi mili — path check karo",
        "KeyError":"Dict mein '{key}' nahi — .get() use karo",
        "IndexError":"Index out of range — length check karo",
        "ZeroDivisionError":"Zero se divide nahi hota",
        "RecursionError":"Infinite recursion — base case add karo",
    }

    @classmethod
    def analyze(cls, error_text: str) -> dict:
        result={"error_type":"Unknown","message":"","line":None,"fix":""}
        for et in cls.FIXES:
            if et in error_text: result["error_type"]=et; break
        m=re.search(r"line (\d+)",error_text,re.I)
        if m: result["line"]=int(m.group(1))
        lines=error_text.strip().split("\n")
        if lines: result["message"]=lines[-1].strip()
        tmpl=cls.FIXES.get(result["error_type"],"Error: {msg}")
        mm=re.search(r"No module named '(.+?)'",error_text)
        nm=re.search(r"name '(.+?)' is not defined",error_text)
        km=re.search(r"KeyError: '?(.+?)'?$",error_text.strip())
        fm=re.search(r"No such file.*: '(.+?)'",error_text)
        result["fix"]=tmpl.format(
            module=mm.group(1) if mm else "module",
            name=nm.group(1) if nm else "var",
            key=km.group(1) if km else "key",
            file=fm.group(1) if fm else "file",
            line=result["line"] or "?",
            msg=result["message"][:80])
        return result

    @classmethod
    def check_safe(cls, cmd: str) -> tuple:
        for p in cls.DANGEROUS:
            if re.search(p,cmd,re.I):
                return False,f"⚠️ DANGEROUS: pattern '{p}' detected"
        return True,"OK"

    @classmethod
    def run_code(cls, code: str, timeout: int=30) -> dict:
        safe,reason=cls.check_safe(code)
        if not safe: return {"success":False,"error":reason,"needs_confirm":True}
        try:
            with tempfile.NamedTemporaryFile(mode="w",suffix=".py",
                                             delete=False,encoding="utf-8") as f:
                f.write(code); tmp=f.name
            r=subprocess.run([sys.executable,tmp],capture_output=True,
                              text=True,timeout=timeout,encoding="utf-8",errors="replace")
            os.unlink(tmp)
            return {"success":r.returncode==0,"output":r.stdout[:2000],
                    "error":r.stderr[:1000],
                    "analysis":cls.analyze(r.stderr) if r.stderr else None}
        except subprocess.TimeoutExpired:
            return {"success":False,"error":f"Timeout ({timeout}s)"}
        except Exception as e:
            return {"success":False,"error":str(e)}

class AssignmentGenerator:
    def __init__(self): ASSIGN_DIR.mkdir(parents=True,exist_ok=True)

    def _styles(self):
        base=getSampleStyleSheet(); s={}
        s["title"]=ParagraphStyle("ATitle",fontName="Times-Bold",fontSize=22,
            alignment=TA_CENTER,spaceAfter=16,leading=30,
            textColor=colors.HexColor("#1a1a2e"))
        s["subtitle"]=ParagraphStyle("ASubtitle",fontName="Times-Roman",fontSize=13,
            alignment=TA_CENTER,spaceAfter=8,textColor=colors.HexColor("#444444"))
        s["meta"]=ParagraphStyle("AMeta",fontName="Times-Roman",fontSize=12,
            spaceAfter=6,textColor=colors.HexColor("#333333"))
        s["h1"]=ParagraphStyle("AH1",fontName="Times-Bold",fontSize=14,
            spaceBefore=20,spaceAfter=12,textColor=colors.HexColor("#1a1a2e"),leading=20)
        s["h2"]=ParagraphStyle("AH2",fontName="Times-Bold",fontSize=13,
            spaceBefore=14,spaceAfter=8,textColor=colors.HexColor("#16213e"),leading=18)
        s["h3"]=ParagraphStyle("AH3",fontName="Times-Bold",fontSize=12,
            spaceBefore=10,spaceAfter=6,textColor=colors.HexColor("#0f3460"),leading=17)
        s["body"]=ParagraphStyle("ABody",fontName="Times-Roman",fontSize=12,
            leading=22,spaceAfter=10,alignment=TA_JUSTIFY)
        s["bullet"]=ParagraphStyle("ABullet",fontName="Times-Roman",fontSize=12,
            leading=20,spaceAfter=6,leftIndent=20,alignment=TA_JUSTIFY)
        s["toc"]=ParagraphStyle("ATOC",fontName="Times-Roman",fontSize=12,
            leading=20,spaceAfter=4,leftIndent=10)
        s["toc_sub"]=ParagraphStyle("ATOCSub",fontName="Times-Roman",fontSize=11,
            leading=18,spaceAfter=3,leftIndent=30,textColor=colors.HexColor("#555555"))
        return s

    def _parse(self, text: str) -> list:
        items=[]
        for line in text.split("\n"):
            s=line.strip()
            if not s: continue
            if re.match(r"^\d+\.\s+[A-Z\u0900-\u097F]",s) and len(s)<120:
                items.append((0,"h1",re.sub(r"^\d+\.\s*","",s)))
            elif s.startswith("### "):
                items.append((2,"h3",s[4:].strip()))
            elif s.startswith("## "):
                items.append((1,"h2",s[3:].strip()))
            elif s.startswith("# "):
                items.append((0,"h1",s[2:].strip()))
            elif s.startswith(("- ","• ","* ")):
                items.append((4,"bullet",s[2:].strip()))
            elif re.match(r"^\d+\.\s",s) and len(s)<300:
                items.append((4,"bullet",s))
            elif s.startswith("**") and s.endswith("**") and len(s)<120:
                items.append((1,"h2",s.strip("*").strip()))
            else:
                items.append((3,"body",s))
        return items

    def build_pdf(self, title: str, student_name: str, subject: str,
                  content: str, output_filename: str=None) -> str:
        if not REPORTLAB_OK:
            return "reportlab not installed: pip install reportlab"
        if not output_filename:
            ts=datetime.datetime.now().strftime("%Y%m%d_%H%M%S")
            output_filename=f"assignment_{ts}.pdf"
        out_path=str(ASSIGN_DIR/output_filename)
        styles=self._styles()
        story=[]

        # Cover
        story.append(Spacer(1,3*cm))
        story.append(Paragraph(subject.upper(),styles["subtitle"]))
        story.append(Spacer(1,0.5*cm))
        story.append(Paragraph(title,styles["title"]))
        story.append(Spacer(1,0.3*cm))
        story.append(HRFlowable(width="75%",thickness=2.5,
                                 color=colors.HexColor("#1a1a2e"),spaceAfter=24))
        story.append(Spacer(1,1.5*cm))
        for lbl,val in [("Submitted By",student_name),("Subject",subject),
                         ("Date",datetime.datetime.now().strftime("%B %d, %Y"))]:
            story.append(Paragraph(f"<b>{lbl}:</b>&nbsp;&nbsp;{val}",styles["meta"]))
            story.append(Spacer(1,0.2*cm))
        story.append(PageBreak())

        # Parse content
        parsed=self._parse(content)

        # TOC
        h1_items=[(i,t) for i,(lvl,typ,t) in enumerate(parsed) if typ=="h1"]
        story.append(Paragraph("Table of Contents",styles["h1"]))
        story.append(HRFlowable(width="100%",thickness=1,
                                 color=colors.HexColor("#cccccc"),spaceAfter=12))
        sec_n=0
        for i,(lvl,typ,txt) in enumerate(parsed):
            if typ=="h1":
                sec_n+=1
                safe=txt.replace("&","&amp;").replace("<","&lt;").replace(">","&gt;")
                story.append(Paragraph(f"{sec_n}.&nbsp;&nbsp;{safe}",styles["toc"]))
            elif typ=="h2":
                safe=txt.replace("&","&amp;").replace("<","&lt;").replace(">","&gt;")
                story.append(Paragraph(f"&nbsp;&nbsp;&nbsp;↳ {safe}",styles["toc_sub"]))
        story.append(PageBreak())

        # Content — KEY FIX: each element is separate, NOT crammed together
        sec_n=0
        chunk=[]
        for (lvl,typ,txt) in parsed:
            if not txt.strip(): continue
            safe=txt.replace("&","&amp;").replace("<","&lt;").replace(">","&gt;")

            # Flush any pending body chunk before headings
            if typ in ("h1","h2","h3") and chunk:
                para_txt=" ".join(chunk)
                # Split into separate paragraphs of ~3-4 sentences each
                sents=re.split(r'(?<=[.!?।])\s+',para_txt)
                group=[]
                for sent in sents:
                    group.append(sent)
                    if len(group)>=3 or len(" ".join(group))>500:
                        story.append(Paragraph(" ".join(group),styles["body"]))
                        story.append(Spacer(1,0.08*cm))
                        group=[]
                if group:
                    story.append(Paragraph(" ".join(group),styles["body"]))
                    story.append(Spacer(1,0.08*cm))
                chunk=[]

            if typ=="h1":
                sec_n+=1
                story.append(Spacer(1,0.4*cm))
                story.append(Paragraph(f"{sec_n}.&nbsp;&nbsp;{safe}",styles["h1"]))
                story.append(HRFlowable(width="100%",thickness=0.5,
                                         color=colors.HexColor("#dddddd"),spaceAfter=6))
            elif typ=="h2":
                story.append(Paragraph(safe,styles["h2"]))
            elif typ=="h3":
                story.append(Paragraph(safe,styles["h3"]))
            elif typ=="bullet":
                if chunk:
                    story.append(Paragraph(" ".join(chunk),styles["body"]))
                    story.append(Spacer(1,0.08*cm))
                    chunk=[]
                story.append(Paragraph(f"&#8226;&nbsp;&nbsp;{safe}",styles["bullet"]))
            else:
                # Accumulate body text, flush by sentence groups
                chunk.append(safe)
                if len(" ".join(chunk))>600:
                    sents=re.split(r'(?<=[.!?।])\s+'," ".join(chunk))
                    group=[]
                    for sent in sents[:-1]:
                        group.append(sent)
                        if len(group)>=3 or len(" ".join(group))>500:
                            story.append(Paragraph(" ".join(group),styles["body"]))
                            story.append(Spacer(1,0.08*cm))
                            group=[]
                    if group:
                        story.append(Paragraph(" ".join(group),styles["body"]))
                        story.append(Spacer(1,0.08*cm))
                    chunk=[sents[-1]] if sents else []

        # Flush remaining
        if chunk:
            story.append(Paragraph(" ".join(chunk),styles["body"]))

        # References
        story.append(PageBreak())
        story.append(Paragraph("References",styles["h1"]))
        story.append(HRFlowable(width="100%",thickness=1,
                                 color=colors.HexColor("#cccccc"),spaceAfter=16))
        story.append(Paragraph(
            "1. Relevant academic textbooks and research papers on this topic.",
            styles["body"]))
        story.append(Paragraph(
            "2. Online resources, journals and publications as referenced in content.",
            styles["body"]))

        def on_page(canvas,doc):
            canvas.saveState()
            canvas.setFont("Times-Roman",9)
            canvas.setFillColor(colors.HexColor("#888888"))
            canvas.drawString(2.5*cm,A4[1]-1.3*cm,subject[:50])
            canvas.drawRightString(A4[0]-2.5*cm,A4[1]-1.3*cm,
                                   datetime.datetime.now().strftime("%Y"))
            canvas.line(2.5*cm,A4[1]-1.5*cm,A4[0]-2.5*cm,A4[1]-1.5*cm)
            canvas.line(2.5*cm,2*cm,A4[0]-2.5*cm,2*cm)
            canvas.drawCentredString(A4[0]/2,1.4*cm,f"Page {doc.page}")
            canvas.drawString(2.5*cm,1.4*cm,title[:50])
            canvas.restoreState()

        try:
            doc=SimpleDocTemplate(out_path,pagesize=A4,
                rightMargin=2.5*cm,leftMargin=2.5*cm,
                topMargin=3*cm,bottomMargin=2.8*cm)
            doc.build(story,onFirstPage=on_page,onLaterPages=on_page)
            print(f"[Assignment] Done: {out_path}")
            return out_path
        except Exception as e:
            print(f"[Assignment] Error: {e}")
            return f"Error: {e}"

class FaceProfiles:
    _profiles: dict={}; _current: Optional[str]=None
    @classmethod
    def load(cls):
        try:
            if PROFILES_FILE.exists():
                with open(PROFILES_FILE) as f: cls._profiles=json.load(f)
        except Exception: cls._profiles={}
    @classmethod
    def get_current_user(cls) -> str: return cls._current or "Unknown"

ADVANCED_FEATURES_TOOLS = [
    {"name":"save_memory","description":"Save note/fact/event to TITAN long-term memory",
     "parameters":{"type":"object","properties":{
         "type":{"type":"string","enum":["note","fact","event","preference"]},
         "title":{"type":"string"},"content":{"type":"string"},
         "tags":{"type":"string"}},"required":["type","content"]}},
    {"name":"search_memory","description":"Search TITAN memory",
     "parameters":{"type":"object","properties":{
         "query":{"type":"string"},"limit":{"type":"integer"}},"required":["query"]}},
    {"name":"get_memory_summary","description":"Get summary of TITAN memories",
     "parameters":{"type":"object","properties":{}}},
    {"name":"take_webcam_photo",
     "description":"Take webcam photo and analyze. Use for: 'meri photo lo', 'kya dikh raha hai', 'mere haath mein kya hai', 'outfit kaisa lag raha hai'",
     "parameters":{"type":"object","properties":{
         "analyze_prompt":{"type":"string","description":"What to analyze in the photo"}}}},
    {"name":"analyze_screen",
     "description":"Take screenshot and analyze screen. Use for: 'meri screen dekho', 'is error ko fix karo', 'screen par kya hai', 'active window dikhao'",
     "parameters":{"type":"object","properties":{
         "analyze_prompt":{"type":"string","description":"What to look for on screen"}}}},
    {"name":"read_file",
     "description":"Read PDF/image/text file with OCR. Searches all folders. Use for: 'is PDF ko padho', 'file read karo', 'question image se assignment banao'",
     "parameters":{"type":"object","properties":{
         "file_path":{"type":"string","description":"Full path OR filename to search for"},
         "action":{"type":"string","enum":["read","summarize","quiz","ocr"]}},"required":["action"]}},
    {"name":"browse_files",
     "description":"Browse files in any folder or find files anywhere on computer",
     "parameters":{"type":"object","properties":{
         "action":{"type":"string","enum":["list_folder","find_file","recent_files","find_pdfs"]},
         "path":{"type":"string"},"filename":{"type":"string"}},"required":["action"]}},
    {"name":"clipboard_action","description":"Read/write clipboard, copy-paste between apps, run in VSCode",
     "parameters":{"type":"object","properties":{
         "action":{"type":"string","enum":["get","set","copy_from_app","paste_to_app","run_in_vscode"]},
         "text":{"type":"string"},"source_app":{"type":"string"},"dest_app":{"type":"string"}},
         "required":["action"]}},
    {"name":"run_code","description":"Run Python code safely, detect errors, suggest fixes",
     "parameters":{"type":"object","properties":{
         "code":{"type":"string"},"timeout":{"type":"integer"}},"required":["code"]}},
    {"name":"analyze_error","description":"Analyze code/terminal error and suggest fix",
     "parameters":{"type":"object","properties":{
         "error_text":{"type":"string"}},"required":["error_text"]}},
    {"name":"screen_control","description":"Control mouse/keyboard: click, type, hotkey, scroll",
     "parameters":{"type":"object","properties":{
         "action":{"type":"string","enum":["click","type","hotkey","scroll","move"]},
         "x":{"type":"integer"},"y":{"type":"integer"},
         "text":{"type":"string"},"keys":{"type":"string"},
         "direction":{"type":"string"},"amount":{"type":"integer"},
         "double":{"type":"boolean"}},"required":["action"]}},
    {"name":"create_assignment_pdf",
     "description":"Create professional assignment PDF 30-40 pages, Times New Roman. Call when user says 'assignment banao', 'PDF banao'. For question images/PDFs: first call read_file to extract questions, then write detailed content, then call this.",
     "parameters":{"type":"object","properties":{
         "title":{"type":"string"},"student_name":{"type":"string"},
         "subject":{"type":"string"},
         "content":{"type":"string","description":"Full markdown content. Use # for main sections (8-10 minimum), ## for sub-sections, normal text for body (200-400 words each section), - for bullets. Write each point as FULL sentence/paragraph NOT single line."},
         "output_filename":{"type":"string"}},"required":["title","content"]}},
]

_cam_result={}; _cam_event=threading.Event(); _cam_threads=[]

def _cam_ok(path):
    _cam_result["path"]=path; _cam_result["error"]=None; _cam_event.set()
def _cam_err(err):
    _cam_result["error"]=err; _cam_result["path"]=None; _cam_event.set()

async def dispatch_advanced_feature(name: str, args: dict) -> str:
    import asyncio

    if name=="save_memory":
        t=args.get("type","note"); c=args.get("content","")
        ti=args.get("title",c[:50])
        tags=[x.strip() for x in args.get("tags","").split(",") if x.strip()]
        if t=="note": TITANMemory.add_note(ti,c,tags)
        else: TITANMemory.add_fact(c)
        return f"✅ Memory saved — {t}: '{ti}'"

    if name=="search_memory":
        results=TITANMemory.search(args.get("query",""),args.get("limit",5))
        if not results: return "Koi memory nahi mili."
        lines=[f"🧠 {len(results)} memories found:"]
        for r in results:
            content=r.get("content") or r.get("fact") or r.get("title") or ""
            lines.append(f"  [{r['type']}] {content[:180]}")
        return "\n".join(lines)

    if name=="get_memory_summary":
        s=TITANMemory.get_summary(); ctx=TITANMemory.get_recent_context(8)
        return (f"🧠 Memory:\n  Conversations:{s['conversations']}\n"
                f"  Notes:{s['notes']}\n  Facts:{s['facts']}\n\n"
                f"Recent:\n{ctx or '(none yet)'}")

    if name=="take_webcam_photo":
        prompt=args.get("analyze_prompt",
            "Describe in detail what you see. Identify objects, people, colors, text, background.")
        _cam_result.clear(); _cam_event.clear()
        t=WebcamThread(CAM_DIR); _cam_threads.append(t)
        t.sig_ok.connect(_cam_ok); t.sig_error.connect(_cam_err)
        t.finished.connect(lambda: _cam_threads.remove(t) if t in _cam_threads else None)
        t.start()
        for _ in range(75):
            await asyncio.sleep(0.2)
            if _cam_event.is_set(): break
        if _cam_result.get("error"):
            return f"❌ Webcam error: {_cam_result['error']}"
        path=_cam_result.get("path")
        if not path or not os.path.exists(path):
            return "❌ Photo capture failed"
        analysis=_analyze_image_with_gemini(path,prompt)
        return f"📸 Photo captured: {Path(path).name}\n\n🔍 Analysis:\n{analysis}\n\n📁 Saved: {path}"

    if name=="analyze_screen":
        prompt=args.get("analyze_prompt",
            "Describe what is visible on screen. Identify open applications, errors, code, text. Be specific and helpful.")
        path=ScreenCapture.capture()
        if not path:
            return "❌ Screenshot failed. Install: pip install pyautogui OR pip install Pillow"
        analysis=_analyze_image_with_gemini(path,prompt)
        return f"🖥️ Screen analyzed\n\n👁️ What I see:\n{analysis}"

    if name=="read_file":
        action=args.get("action","read")
        file_path=args.get("file_path","")
        if not file_path:
            return "📂 File path batao. Bol: '[filename] padho' ya full path do."

        ext=Path(file_path).suffix.lower()

        # Word Files (.docx, .doc)
        if ext in {".docx", ".doc"}:
            actual, text = DocReader.find_and_extract(file_path)
            if not actual: return text
            if action=="read": return f"📄 Word File ({Path(actual).name}):\n\n{text[:4000]}"
            if action=="summarize": return f"📝 Summary:\n{PDFReader.summarize(text)}" # Reuse summarizer
            return text[:3000]

        # Image → OCR
        if ext in {".png",".jpg",".jpeg",".bmp",".gif",".tiff",".webp"}:
            actual=file_path
            if not os.path.exists(actual):
                found=FileBrowser.find_file(file_path)
                if found: actual=found[0]
                else: return f"❌ Image not found: {file_path}"
            text=extract_text_from_image(actual)
            return f"📄 Text extracted from image:\n\n{text}"
        # PDF
        actual,text=PDFReader.find_and_extract(file_path)
        if not actual: return text
        if action=="read": return f"📄 PDF ({Path(actual).name}):\n\n{text[:4000]}"
        if action=="summarize": return f"📝 Summary:\n{PDFReader.summarize(text)}"
        if action=="quiz":
            q=PDFReader.quiz(text)
            if not q: return "❌ Could not generate quiz"
            lines=[f"📚 Quiz ({len(q)} questions):"]
            for qi in q: lines.append(f"\n{qi['q']}\nAnswer: {qi['answer']}")
            return "\n".join(lines)
        if action=="ocr": return f"📄 Content:\n{text[:3000]}"
        return text[:2000]

    if name=="browse_files":
        action=args.get("action","recent_files")
        if action=="recent_files": return FileBrowser.recent_files()
        if action=="find_pdfs":
            pdfs=FileBrowser.find_pdfs()
            if not pdfs: return "No PDFs found"
            return "📄 PDFs:\n"+"\n".join(f"  • {p}" for p in pdfs[:15])
        if action=="find_file":
            fn=args.get("filename","")
            if not fn: return "❌ Filename batao"
            found=FileBrowser.find_file(fn)
            if not found: return f"❌ '{fn}' nahi mila"
            return f"🔍 Found:\n"+"\n".join(f"  • {p}" for p in found)
        if action=="list_folder":
            path=args.get("path",str(Path.home()))
            return FileBrowser.list_folder(path)
        return "Unknown action"

    if name=="clipboard_action":
        action=args.get("action","")
        if action=="get":
            t=ClipboardAutomation.get()
            return f"📋 Clipboard:\n{t[:500]}" if t else "Clipboard empty"
        if action=="set":
            ok=ClipboardAutomation.set(args.get("text",""))
            return "✅ Clipboard set" if ok else "❌ Failed"
        if action=="copy_from_app":
            t=ClipboardAutomation.copy_from_app(args.get("source_app",""))
            return f"📋 Copied:\n{t[:500]}"
        if action=="paste_to_app":
            ok=ClipboardAutomation.paste_to_app(args.get("dest_app",""),args.get("text",""))
            return "✅ Pasted" if ok else "❌ Failed"
        if action=="run_in_vscode":
            r=ClipboardAutomation.run_in_vscode(args.get("text",ClipboardAutomation.get()))
            return r.get("message","Done") if r["success"] else f"❌ {r['error']}"
        return "Unknown action"

    if name=="run_code":
        r=ErrorDetector.run_code(args.get("code",""),args.get("timeout",30))
        if r.get("needs_confirm"): return f"⚠️ Safety check: {r['error']}"
        if r["success"]: return f"✅ Code successful:\n{r.get('output','(no output)')}"
        a=r.get("analysis") or {}
        return f"❌ Error:\n{r.get('error','')[:400]}" + (f"\n💡 Fix: {a['fix']}" if a.get("fix") else "")

    if name=="analyze_error":
        a=ErrorDetector.analyze(args.get("error_text",""))
        return (f"🔍 Error: {a['error_type']}\n  Line: {a.get('line','?')}\n"
                f"  Message: {a['message']}\n  💡 Fix: {a['fix']}")

    if name=="screen_control":
        if not PYAUTOGUI_OK: return "❌ pyautogui not installed: pip install pyautogui"
        action=args.get("action","")
        try:
            if action=="click":
                x,y=args.get("x",0),args.get("y",0)
                (pyautogui.doubleClick if args.get("double") else pyautogui.click)(x,y)
                return f"✅ Clicked ({x},{y})"
            if action=="type":
                pyautogui.typewrite(args.get("text",""),interval=0.03)
                return "✅ Typed"
            if action=="hotkey":
                keys=[k.strip() for k in args.get("keys","").split(",")]
                pyautogui.hotkey(*keys); return f"✅ Hotkey: {'+'.join(keys)}"
            if action=="scroll":
                d=-args.get("amount",3) if args.get("direction","down")=="down" else args.get("amount",3)
                pyautogui.scroll(d); return f"✅ Scrolled"
            if action=="move":
                pyautogui.moveTo(args.get("x",0),args.get("y",0)); return "✅ Moved"
        except Exception as e: return f"❌ Error: {e}"
        return "Unknown action"

    if name=="create_assignment_pdf":
        gen=AssignmentGenerator()
        title=args.get("title","Assignment")
        content=args.get("content","")
        student=args.get("student_name","Student")
        subject=args.get("subject","General Studies")
        out_file=args.get("output_filename")
        if not content or len(content)<200:
            return "❌ Content too short. Write full detailed content first (8+ sections, 200+ words each)."
        path=gen.build_pdf(title,student,subject,content,out_file)
        if path and os.path.exists(path):
            kb=os.path.getsize(path)//1024
            return (f"✅ Assignment PDF ready!\n📁 Path: {path}\n📊 Size: {kb} KB\n"
                    f"Heading: 14pt Times-Bold ✓\nBody: 12pt Times-Roman justified ✓")
        return f"❌ PDF failed: {path}"

    return f"Unknown tool: {name}"

class TITANFeaturesPanel(QFrame):
    sig_action=pyqtSignal(str)
    def __init__(self,parent=None):
        super().__init__(parent)
        self.setStyleSheet(f"QFrame{{background:transparent;border:1px solid rgba(0,212,255,0.1);border-radius:8px;}}"
                           f"QLabel{{background:transparent;color:{TEXT};}}"
                           f"QPushButton{{background:rgba(0,212,255,0.06);border:1px solid rgba(0,212,255,0.18);"
                           f"border-radius:5px;color:{CYAN};font-size:9px;font-family:'Courier New';padding:3px 5px;}}"
                           f"QPushButton:hover{{background:rgba(0,212,255,0.14);}}")
        self._build(); TITANMemory.load(); FaceProfiles.load(); self._refresh()
        QTimer(self,timeout=self._refresh,interval=30000).start()

    def _build(self):
        vl=QVBoxLayout(self); vl.setContentsMargins(8,8,8,8); vl.setSpacing(5)
        hdr=QLabel("⚡ FEATURES")
        hdr.setStyleSheet(f"color:{CYAN};font-size:9px;font-family:'Courier New';letter-spacing:2px;font-weight:bold;")
        vl.addWidget(hdr)
        self._mem_lbl=QLabel("📦 Memory: loading...")
        self._mem_lbl.setStyleSheet(f"color:{MUTED};font-size:9px;font-family:'Courier New';")
        vl.addWidget(self._mem_lbl)
        grid=QGridLayout(); grid.setSpacing(3)
        btns=[("📸 Webcam","Take a webcam photo and tell me what you see"),
               ("🖥️ Screen","Look at my screen and describe what is happening"),
               ("📋 Clipboard","What is in my clipboard right now?"),
               ("🧠 Memory","What do you remember about me so far?"),
               ("📄 Find PDFs","Find all PDF files on my computer"),
               ("📝 Assignment","I want to create an assignment PDF")]
        for i,(lbl,act) in enumerate(btns):
            btn=QPushButton(lbl); btn.setFixedHeight(22)
            btn.clicked.connect(lambda _,a=act: self.sig_action.emit(a))
            grid.addWidget(btn,i//2,i%2)
        vl.addLayout(grid)

    def _refresh(self):
        s=TITANMemory.get_summary(); total=sum(s.values())
        self._mem_lbl.setText(f"📦 {total} memories ({s['conversations']} chats, {s['notes']} notes)")

class AssignmentDialog(QDialog):
    sig_create=pyqtSignal(dict)
    def __init__(self,parent=None):
        super().__init__(parent)
        self.setWindowTitle("TITAN — Assignment Generator")
        self.setMinimumSize(540,440)
        self.setStyleSheet(f"QDialog,QWidget{{background:{BG};color:{TEXT};font-family:'Segoe UI';}}"
                           f"QLineEdit,QTextEdit{{background:rgba(255,255,255,0.05);border:1px solid rgba(0,212,255,0.22);"
                           f"border-radius:6px;color:{TEXT};padding:6px;font-size:12px;}}"
                           f"QPushButton{{background:rgba(232,64,106,0.14);border:1px solid {ACCENT};"
                           f"border-radius:8px;color:{ACCENT2};padding:8px 16px;font-size:11px;font-weight:bold;}}"
                           f"QPushButton:hover{{background:rgba(232,64,106,0.26);}}"
                           f"QLabel{{color:{MUTED};font-size:11px;}}")
        self._pdf_path=None; self._img_path=None; self._build()

    def _build(self):
        vl=QVBoxLayout(self); vl.setSpacing(10); vl.setContentsMargins(20,18,20,18)
        vl.addWidget(QLabel("📝 ASSIGNMENT GENERATOR",styleSheet=f"color:{ACCENT2};font-size:14px;font-weight:bold;"))
        vl.addWidget(QLabel("Topic / Title:"))
        self.topic=QLineEdit(); self.topic.setPlaceholderText("e.g., AI in Education"); vl.addWidget(self.topic)
        row1=QHBoxLayout()
        row1.addWidget(QLabel("Student Name:")); self.name=QLineEdit(); self.name.setPlaceholderText("Your name"); row1.addWidget(self.name)
        row1.addWidget(QLabel("Subject:")); self.sub=QLineEdit(); self.sub.setPlaceholderText("Subject"); row1.addWidget(self.sub)
        vl.addLayout(row1)
        vl.addWidget(QLabel("Question Image / PDF (optional — TITAN will read questions):"))
        src_row=QHBoxLayout()
        self._src_lbl=QLabel("No file selected"); self._src_lbl.setStyleSheet(f"color:{MUTED};font-size:10px;"); src_row.addWidget(self._src_lbl)
        img_btn=QPushButton("📷 Image"); img_btn.setFixedWidth(90)
        img_btn.setStyleSheet(f"background:rgba(0,212,255,0.1);border:1px solid {CYAN};color:{CYAN};border-radius:7px;padding:5px;font-size:10px;")
        img_btn.clicked.connect(_ALIRA_slot(self._pick_image)); src_row.addWidget(img_btn)
        pdf_btn=QPushButton("📄 PDF"); pdf_btn.setFixedWidth(80)
        pdf_btn.setStyleSheet(f"background:rgba(255,208,96,0.1);border:1px solid {GOLD};color:{GOLD};border-radius:7px;padding:5px;font-size:10px;")
        pdf_btn.clicked.connect(_ALIRA_slot(self._pick_pdf)); src_row.addWidget(pdf_btn)
        vl.addLayout(src_row)
        vl.addWidget(QLabel("Additional Instructions:"))
        self.inst=QTextEdit(); self.inst.setPlaceholderText("e.g., 'Include case studies', '35+ pages', 'Focus on Indian context'")
        self.inst.setMaximumHeight(70); vl.addWidget(self.inst)
        btns=QHBoxLayout()
        cancel=QPushButton("Cancel")
        cancel.setStyleSheet(f"background:rgba(80,80,100,0.12);border:1px solid {MUTED};color:{MUTED};border-radius:8px;padding:8px 14px;")
        cancel.clicked.connect(_ALIRA_slot(self.reject))
        go=QPushButton("🚀 Generate Assignment PDF"); go.clicked.connect(_ALIRA_slot(self._go))
        btns.addWidget(cancel); btns.addStretch(); btns.addWidget(go); vl.addLayout(btns)

    def _pick_image(self):
        p,_=QFileDialog.getOpenFileName(self,"Select Question Image","","Images (*.png *.jpg *.jpeg *.bmp *.tiff *.webp)")
        if p:
            self._img_path=p; self._src_lbl.setText(f"📷 {Path(p).name}")
            self._src_lbl.setStyleSheet(f"color:{CYAN};font-size:10px;")

    def _pick_pdf(self):
        p,_=QFileDialog.getOpenFileName(self,"Select Question PDF","","PDF (*.pdf)")
        if p:
            self._pdf_path=p; self._src_lbl.setText(f"📄 {Path(p).name}")
            self._src_lbl.setStyleSheet(f"color:{GOLD};font-size:10px;")

    def _go(self):
        if not self.topic.text().strip():
            QMessageBox.warning(self,"Input Required","Topic dena padega"); return
        self.sig_create.emit({"topic":self.topic.text().strip(),
            "student_name":self.name.text().strip() or "Student",
            "subject":self.sub.text().strip() or "General Studies",
            "instructions":self.inst.toPlainText().strip(),
            "image_path":self._img_path,"pdf_path":self._pdf_path})
        self.accept()

def init_advanced_features(api_key: str=None):
    global _API_KEY
    if api_key: _API_KEY=api_key
    TITANMemory.load(); FaceProfiles.load()
    print(f"[TITANFeatures v2] Root:{ROOT_PATH}")
    print(f"[TITANFeatures v2] CV2:{CV2_OK} PyAutoGUI:{PYAUTOGUI_OK} OCR:{OCR_OK} RL:{REPORTLAB_OK}")
